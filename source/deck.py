"""
Aurora Dark Premium — moteur de rendu double (PPTX + aperçu PNG).
Un même "spec" de slide alimente python-pptx et un rendu PIL de vérification.
"""
import copy
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
import math, os

EMU = 914400
def IN(v): return Emu(int(v*EMU))
PXIN = 100  # preview px per inch
W_IN, H_IN = 13.3333, 7.5

FONTDIR = "/usr/share/fonts/truetype/custom"
FONT_FILES = {
    "Montserrat": "Montserrat-Regular.ttf",
    "Montserrat Medium": "Montserrat-Medium.ttf",
    "Montserrat SemiBold": "Montserrat-SemiBold.ttf",
    "Montserrat Bold": "Montserrat-Bold.ttf",
    "Montserrat ExtraBold": "Montserrat-ExtraBold.ttf",
    "Montserrat Black": "Montserrat-Black.ttf",
    "Inter": "Inter-Regular.ttf",
    "Inter Medium": "Inter-Medium.ttf",
    "Inter SemiBold": "Inter-SemiBold.ttf",
    "Inter Bold": "Inter-Bold.ttf",
}

# ---------- palette ----------
INK        = "0A0E1A"
INK2       = "0E1426"
INK_BOT    = "070A12"
CARD       = "151D33"
CARD2      = "1A2440"
CARD_HI    = "1E2A4A"
BORDER     = "263352"
BORDER_HI  = "3A4A78"
WHITE      = "FFFFFF"
TXT        = "F4F7FF"
TXT2       = "AEB9D6"
TXT3       = "7E8AAC"
# accents
BLUE       = "5B8DEF"
VIOLET     = "9B6DFF"
GOLD       = "F6B43C"
GOLD2      = "FFD36E"
CYAN       = "31D7E6"
ROSE       = "FB7185"
GREEN      = "3DDC97"

GR_PRIMARY = (BLUE, VIOLET)
GR_GOLD    = (GOLD, GOLD2)
GR_CYAN    = (CYAN, BLUE)
GR_ROSE    = (ROSE, VIOLET)
GR_GREEN   = (GREEN, CYAN)

def hx(c): return tuple(int(c[i:i+2],16) for i in (0,2,4))

# ===================================================================
#  SPEC builders — chaque slide est une liste d'éléments (dicts)
# ===================================================================
def el(t, **k): k["t"]=t; return k

# ===================================================================
#  PPTX RENDERER
# ===================================================================
def _set_fill(shape, color):
    shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(color)

def _no_line(shape):
    shape.line.fill.background()

def _grad_xml(angle_deg, stops):
    # stops: list of (pos0-100, hexcolor, alpha0-100)
    ang = int(angle_deg*60000)
    s=''
    for pos,col,al in stops:
        a = f'<a:alpha val="{int(al*1000)}"/>' if al<100 else ''
        s+=f'<a:gs pos="{int(pos*1000)}"><a:srgbClr val="{col}">{a}</a:srgbClr></a:gs>'
    return (f'<a:gradFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
            f'rotWithShape="1"><a:gsLst>{s}</a:gsLst>'
            f'<a:lin ang="{ang}" scaled="1"/></a:gradFill>')

def _radial_xml(stops):
    s=''
    for pos,col,al in stops:
        a=f'<a:alpha val="{int(al*1000)}"/>'
        s+=f'<a:gs pos="{int(pos*1000)}"><a:srgbClr val="{col}">{a}</a:srgbClr></a:gs>'
    return (f'<a:gradFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
            f'rotWithShape="1"><a:gsLst>{s}</a:gsLst>'
            f'<a:path path="circle"><a:fillToRect l="50000" t="50000" r="50000" b="50000"/></a:path>'
            f'</a:gradFill>')

def _apply_grad(shape, xml):
    spPr = shape.fill._xPr  # spPr
    # remove existing fill nodes
    for tag in ('a:noFill','a:solidFill','a:gradFill','a:blipFill','a:pattFill','a:grpFill'):
        for e in spPr.findall(qn(tag)): spPr.remove(e)
    from lxml import etree
    node = etree.fromstring(xml)
    # insert after prstGeom/custGeom, before line(a:ln)
    ln = spPr.find(qn('a:ln'))
    if ln is not None: ln.addprevious(node)
    else: spPr.append(node)

def _shadow(shape, blur=18, dist=10, dir=5400000, alpha=58, color="000000"):
    spPr = shape._element.spPr
    from lxml import etree
    xml=(f'<a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
         f'<a:outerShdw blurRad="{blur*12700}" dist="{dist*12700}" dir="{dir}" rotWithShape="0">'
         f'<a:srgbClr val="{color}"><a:alpha val="{alpha*1000}"/></a:srgbClr></a:outerShdw></a:effectLst>')
    spPr.append(etree.fromstring(xml))

def _set_spc(run, spc_pt):
    run._r.get_or_add_rPr().set('spc', str(int(spc_pt*100)))

def _run_grad(run, angle, stops):
    rPr = run._r.get_or_add_rPr()
    for tag in ('a:solidFill','a:gradFill'):
        for e in rPr.findall(qn(tag)): rPr.remove(e)
    from lxml import etree
    node = etree.fromstring(_grad_xml(angle, stops))
    # CT_TextCharacterProperties order: ln, fill, effectLst, ..., latin, ea, cs, sym...
    # gradFill (fill) must precede latin/ea/cs/sym/uLn/highlight/effectLst
    anchor=None
    for tag in ('a:effectLst','a:highlight','a:uLnTx','a:uLn','a:uFillTx','a:uFill',
                'a:latin','a:ea','a:cs','a:sym','a:hlinkClick','a:hlinkMouseOver','a:rtl','a:extLst'):
        f=rPr.find(qn(tag))
        if f is not None: anchor=f; break
    ln=rPr.find(qn('a:ln'))
    if anchor is not None: anchor.addprevious(node)
    elif ln is not None: ln.addnext(node)
    else: rPr.append(node)

def render_pptx(slides, out):
    prs = Presentation(); prs.slide_width=IN(W_IN); prs.slide_height=IN(H_IN)
    blank = prs.slide_layouts[6]
    for spec in slides:
        s = prs.slides.add_slide(blank)
        for e in spec:
            _pptx_el(s, e)
    prs.save(out)

def _add_rrect(s, e):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, IN(e["x"]),IN(e["y"]),IN(e["w"]),IN(e["h"]))
    rad = e.get("radius",0.14)
    try: shp.adjustments[0]= rad / min(e["w"],e["h"])
    except: pass
    if e.get("grad"):
        _apply_grad(shp, _grad_xml(e.get("angle",90), e["grad"]))
    else:
        _set_fill(shp, e.get("fill",CARD))
    if e.get("line"):
        shp.line.color.rgb=RGBColor.from_string(e["line"]); shp.line.width=Pt(e.get("line_w",1))
    else:
        _no_line(shp)
    if e.get("shadow"): _shadow(shp, **(e["shadow"] if isinstance(e["shadow"],dict) else {}))
    return shp

def _pptx_el(s, e):
    t=e["t"]
    if t=="bg":
        shp=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,IN(W_IN),IN(H_IN))
        _apply_grad(shp,_grad_xml(e.get("angle",105),e["grad"])); _no_line(shp)
    elif t=="glow":
        d=e["r"]*2
        shp=s.shapes.add_shape(MSO_SHAPE.OVAL,IN(e["cx"]-e["r"]),IN(e["cy"]-e["r"]),IN(d),IN(d))
        _apply_grad(shp,_radial_xml([(0,e["color"],e.get("alpha",55)),(100,e["color"],0)])); _no_line(shp)
    elif t=="rrect":
        _add_rrect(s,e)
    elif t=="oval":
        shp=s.shapes.add_shape(MSO_SHAPE.OVAL,IN(e["x"]),IN(e["y"]),IN(e["w"]),IN(e["h"]))
        if e.get("grad"): _apply_grad(shp,_grad_xml(e.get("angle",90),e["grad"]))
        else: _set_fill(shp,e.get("fill",CARD))
        if e.get("line"): shp.line.color.rgb=RGBColor.from_string(e["line"]); shp.line.width=Pt(e.get("line_w",1))
        else: _no_line(shp)
        if e.get("shadow"): _shadow(shp)
    elif t=="line":
        shp=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,IN(e["x"]),IN(e["y"]),IN(e["w"]),IN(e["h"]))
        if e.get("grad"): _apply_grad(shp,_grad_xml(e.get("angle",0),e["grad"]))
        else: _set_fill(shp,e.get("fill",BORDER))
        _no_line(shp)
    elif t=="icon":
        s.shapes.add_picture(e["file"],IN(e["x"]),IN(e["y"]),IN(e["w"]),IN(e["h"]))
    elif t=="text":
        _pptx_text(s,e)

def _pptx_text(s,e):
    tb=s.shapes.add_textbox(IN(e["x"]),IN(e["y"]),IN(e["w"]),IN(e["h"]))
    tf=tb.text_frame; tf.word_wrap=True
    tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
    va=e.get("valign","top")
    tf.vertical_anchor={"top":MSO_ANCHOR.TOP,"middle":MSO_ANCHOR.MIDDLE,"bottom":MSO_ANCHOR.BOTTOM}[va]
    paras=e["paras"]
    for i,p in enumerate(paras):
        para = tf.paragraphs[0] if i==0 else tf.add_paragraph()
        para.alignment={"l":PP_ALIGN.LEFT,"c":PP_ALIGN.CENTER,"r":PP_ALIGN.RIGHT}[p.get("align","l")]
        if p.get("space_before") is not None: para.space_before=Pt(p["space_before"])
        if p.get("space_after") is not None: para.space_after=Pt(p["space_after"])
        if p.get("line"):
            para.line_spacing=p["line"]
        for r in p["runs"]:
            run=para.add_run(); run.text=r["t"]
            f=run.font; f.name=r.get("font","Inter"); f.size=Pt(r["size"])
            f.bold=r.get("bold",False)
            if r.get("spc"): _set_spc(run,r["spc"])
            if r.get("grad"): _run_grad(run,r.get("gangle",0),r["grad"])
            else: f.color.rgb=RGBColor.from_string(r.get("color",TXT))

# ===================================================================
#  PIL PREVIEW RENDERER
# ===================================================================
_fcache={}
def _pil_font(name,size):
    key=(name,size)
    if key not in _fcache:
        _fcache[key]=ImageFont.truetype(os.path.join(FONTDIR,FONT_FILES.get(name,"Inter-Regular.ttf")),size)
    return _fcache[key]

def make_circle_avatar(src, out, d=360):
    """Recadre une photo en cercle (fond transparent) pour servir d'avatar."""
    im=Image.open(src).convert("RGBA")
    w,h=im.size; sq=min(w,h)
    im=im.crop(((w-sq)//2,(h-sq)//2,(w-sq)//2+sq,(h-sq)//2+sq)).resize((d,d),Image.LANCZOS)
    mask=Image.new("L",(d,d),0); ImageDraw.Draw(mask).ellipse([0,0,d-1,d-1],fill=255)
    res=Image.new("RGBA",(d,d),(0,0,0,0)); res.paste(im,(0,0),mask)
    res.save(out); return out

def _lerp(a,b,t): return tuple(int(a[i]+(b[i]-a[i])*t) for i in range(3))

def _grad_img(w,h,angle,stops):
    # linear gradient approx (stops: pos,hex,alpha)
    img=Image.new("RGBA",(w,h))
    px=img.load()
    rad=math.radians(angle)
    dx,dy=math.cos(rad),math.sin(rad)
    cols=[(p/100.0,hx(c),al/100.0) for p,c,al in stops]
    for y in range(h):
        for x in range(w):
            proj=((x/w)*dx+(y/h)*dy+1)/2
            proj=min(1,max(0,proj))
            # find segment
            c0=cols[0]; c1=cols[-1]
            for i in range(len(cols)-1):
                if cols[i][0]<=proj<=cols[i+1][0]:
                    c0,c1=cols[i],cols[i+1]; break
            span=(c1[0]-c0[0]) or 1; tt=(proj-c0[0])/span
            col=_lerp(c0[1],c1[1],tt); al=int((c0[2]+(c1[2]-c0[2])*tt)*255)
            px[x,y]=(col[0],col[1],col[2],al)
    return img

def _radial_img(d,color,alpha):
    img=Image.new("RGBA",(d,d),(0,0,0,0))
    px=img.load(); c=hx(color); r=d/2
    for y in range(d):
        for x in range(d):
            dist=math.hypot(x-r,y-r)/r
            if dist>1: a=0
            else: a=int(alpha/100*255*(1-dist)**1.6)
            px[x,y]=(c[0],c[1],c[2],a)
    return img

def _wrap(draw,text,font,maxw,spc=0):
    words=text.split(" "); lines=[]; cur=""
    def width(s):
        w=draw.textlength(s,font=font)
        if spc: w+=spc*max(0,len(s)-1)
        return w
    for w in words:
        test=(cur+" "+w).strip()
        if width(test)<=maxw or not cur: cur=test
        else: lines.append(cur); cur=w
    if cur: lines.append(cur)
    return lines

def _draw_text_run(draw,x,y,s,font,fill,spc=0):
    if not spc:
        draw.text((x,y),s,font=font,fill=fill); return draw.textlength(s,font=font)
    cx=x
    for ch in s:
        draw.text((cx,y),ch,font=font,fill=fill); cx+=draw.textlength(ch,font=font)+spc
    return cx-x

def render_preview(spec, scale=1.0):
    W=int(W_IN*PXIN*scale); H=int(H_IN*PXIN*scale)
    base=Image.new("RGBA",(W,H),(10,14,26,255))
    draw=ImageDraw.Draw(base)
    def S(v): return int(v*PXIN*scale)
    for e in spec:
        t=e["t"]
        if t=="bg":
            g=_grad_img(W,H,e.get("angle",105),e["grad"]); base.alpha_composite(g)
            draw=ImageDraw.Draw(base)
        elif t=="glow":
            d=S(e["r"]*2)
            if d>0:
                g=_radial_img(d,e["color"],e.get("alpha",55))
                base.alpha_composite(g,(S(e["cx"]-e["r"]),S(e["cy"]-e["r"])))
        elif t in ("rrect","oval","line"):
            x,y,w,h=S(e["x"]),S(e["y"]),max(1,S(e["w"])),max(1,S(e["h"]))
            rad=min(S(e.get("radius",0.14)), (min(w,h)-1)//2) if t=="rrect" else (min(w,h)//2 if t=="oval" else 0)
            rad=max(0,rad)
            layer=Image.new("RGBA",(W,H),(0,0,0,0)); ld=ImageDraw.Draw(layer)
            if e.get("grad"):
                gimg=_grad_img(max(1,w),max(1,h),e.get("angle",90),e["grad"])
                mask=Image.new("L",(max(1,w),max(1,h)),0); md=ImageDraw.Draw(mask)
                if t=="oval": md.ellipse([0,0,w-1,h-1],fill=255)
                else: md.rounded_rectangle([0,0,w-1,h-1],radius=rad,fill=255)
                layer.paste(gimg,(x,y),mask)
            else:
                col=hx(e.get("fill",CARD))+(255,)
                if t=="oval": ld.ellipse([x,y,x+w,y+h],fill=col)
                elif t=="line": ld.rectangle([x,y,x+w,y+h],fill=col)
                else: ld.rounded_rectangle([x,y,x+w,y+h],radius=rad,fill=col)
            if e.get("line"):
                lc=hx(e["line"])+(255,)
                if t=="oval": ld.ellipse([x,y,x+w,y+h],outline=lc,width=max(1,S(e.get("line_w",1)/72)))
                else: ld.rounded_rectangle([x,y,x+w,y+h],radius=rad,outline=lc,width=max(1,int(e.get("line_w",1)*scale)))
            base.alpha_composite(layer); draw=ImageDraw.Draw(base)
        elif t=="icon":
            im=Image.open(e["file"]).convert("RGBA").resize((S(e["w"]),S(e["h"])))
            if e.get("tint"):
                c=hx(e["tint"]); solid=Image.new("RGBA",im.size,(c[0],c[1],c[2],0))
                r,g,b,a=im.split(); solid.putalpha(a); im=solid
            base.alpha_composite(im,(S(e["x"]),S(e["y"])))
        elif t=="text":
            _preview_text(base,draw,e,S,scale)
    return base.convert("RGB")

def _preview_text(base,draw,e,S,scale):
    x,y,w,h=S(e["x"]),S(e["y"]),S(e["w"]),S(e["h"])
    fs=scale*1.333
    blocks=[]  # each: (align,linef,sb,sa, line) where line=list of (text,font,r,spc)
    for p in e["paras"]:
        align=p.get("align","l"); linef=p.get("line",1.0)
        sb=p.get("space_before",0); sa=p.get("space_after",0)
        # tokenize across runs preserving run boundaries
        toks=[]
        for r in p["runs"]:
            font=_pil_font(r.get("font","Inter"),max(1,int(r["size"]*fs)))
            spc=int(r.get("spc",0)*fs)
            words=r["t"].split(" ")
            for wi,wd in enumerate(words):
                tt=wd+(" " if wi<len(words)-1 else "")
                if tt: toks.append((tt,font,r,spc))
        # greedy wrap
        def tokw(t):
            tx,f,r,sp=t; ww=draw.textlength(tx,font=f)
            if sp: ww+=sp*max(0,len(tx)-1)
            return ww
        lines=[]; cur=[]; curw=0
        for t in toks:
            tw=tokw(t)
            if cur and curw+tw>w:
                lines.append(cur); cur=[t]; curw=tw
            else:
                cur.append(t); curw+=tw
        if cur: lines.append(cur)
        if not lines: lines=[[]]
        for li,ln in enumerate(lines):
            blocks.append((align,linef,sb if li==0 else 0, sa if li==len(lines)-1 else 0, ln))
    def line_h(ln): return max([f.size for _,f,_,_ in ln], default=int(12*fs))
    total=sum(line_h(ln)*lf + sb*fs + sa*fs for _,lf,sb,sa,ln in blocks)
    va=e.get("valign","top")
    cy = y if va=="top" else (y+(h-total)/2 if va=="middle" else y+h-total)
    for align,linef,sb,sa,ln in blocks:
        cy+=sb*fs
        sh=line_h(ln)
        lw=sum((draw.textlength(s,font=f)+(sp*max(0,len(s)-1) if sp else 0)) for s,f,_,sp in ln)
        if align=="c": sx=x+(w-lw)/2
        elif align=="r": sx=x+w-lw
        else: sx=x
        for s,f,r,sp in ln:
            fill=hx(r.get("color",TXT))
            if r.get("grad"):
                cols=[hx(c) for _,c,_ in r["grad"]]; fill=_lerp(cols[0],cols[-1],0.5)
            st=s.rstrip(" ") if s.endswith(" ") and ln[-1] is (s,f,r,sp) else s
            adv=_draw_text_run(draw,sx,cy,s,f,fill+(255,),sp)
            sx+=adv
        cy+= sh*linef + sa*fs

def _seg_w(draw,s,f,sp):
    w=draw.textlength(s,font=f)
    if sp: w+=sp*max(0,len(s)-1)
    return w
